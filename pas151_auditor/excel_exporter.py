# -*- coding: utf-8 -*-
"""
Exportador de Reporte Excel Multi-Hoja para la Auditoría PAS 151 CONAF.
Genera el archivo `00_Reporte_revision_campos_<PREDIO>.xlsx` con 6 hojas:
1. `Revision_Campos_R`: Tabla idéntica a `resultados_df` del script R (`write_xlsx`).
2. `Resumen_Ejecutivo`: KPIs generales del predio, archivos y estado global.
3. `Validacion_Capas`: Matriz de cumplimiento por shapefile (CRS, 2D, Orden, Dominios, Sup_ha, Coord_XY, Texto, Geometría).
4. `Validaciones_Espaciales`: Resultados de los 7 controles espaciales e inter-capas.
5. `Incidencias_Detalladas`: Listado completo de errores y advertencias con filas afectadas y recomendación.
6. `Plantilla_CONAF_PAS151`: Diccionario de datos y dominios oficiales aplicados.
"""

from pathlib import Path
from typing import Union
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from pas151_auditor.engine import PAS151AuditResult
from pas151_auditor.templates_conaf import DOMINIOS_CONAF


class PAS151ExcelExporter:
    """Genera el libro Excel consolidado de auditoría PAS 151 CONAF."""

    def __init__(self) -> None:
        self.header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
        self.header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")

        self.sub_fill = PatternFill(start_color="2E5B88", end_color="2E5B88", fill_type="solid")
        self.sub_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")

        self.title_font = Font(name="Segoe UI", size=14, bold=True, color="1B365D")
        self.normal_font = Font(name="Segoe UI", size=10)
        self.bold_font = Font(name="Segoe UI", size=10, bold=True)

        self.fill_ok = PatternFill(start_color="D4EFDF", end_color="D4EFDF", fill_type="solid")
        self.fill_err = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")
        self.fill_warn = PatternFill(start_color="FCF3CF", end_color="FCF3CF", fill_type="solid")
        self.fill_na = PatternFill(start_color="F2F4F4", end_color="F2F4F4", fill_type="solid")

        thin = Side(border_style="thin", color="D5D8DC")
        self.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def export(self, result: PAS151AuditResult, output_path: Union[str, Path]) -> Path:
        out_file = Path(output_path)
        out_file.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        self._build_sheet_tabla_r(wb, result)
        self._build_sheet_resumen(wb, result)
        self._build_sheet_validacion_capas(wb, result)
        self._build_sheet_espaciales(wb, result)
        self._build_sheet_incidencias(wb, result)
        self._build_sheet_plantilla(wb, result)

        wb.save(out_file)
        return out_file

    def _build_sheet_tabla_r(self, wb: openpyxl.Workbook, result: PAS151AuditResult) -> None:
        """
        Hoja 1: Replica exactamente la estructura de columnas que genera el script R:
        shape | nombre | Tipo_Plant | Tipo_Shp | Existe_Campo | Coin_Tipo_Campo |
        Long_Plant | Long_Shp | Long_Correcta | Dec_Plant | Dec_Shp | Dec_Correcto | orden_correcto
        """
        ws = wb.create_sheet(title="Revision_Campos_R")
        headers = [
            "shape",
            "nombre",
            "Tipo_Plant",
            "Tipo_Shp",
            "Existe_Campo",
            "Coin_Tipo_Campo",
            "Long_Plant",
            "Long_Shp",
            "Long_Correcta",
            "Dec_Plant",
            "Dec_Shp",
            "Dec_Correcto",
            "orden_correcto",
        ]
        ws.append(headers)
        self._style_header(ws)

        for r in result.filas_tabla_r:
            row_vals = [
                r.shape,
                r.nombre,
                r.Tipo_Plant,
                r.Tipo_Shp if r.Tipo_Shp is not None else "NA",
                r.Existe_Campo,
                r.Coin_Tipo_Campo if r.Coin_Tipo_Campo is not None else "NA",
                r.Long_Plant if r.Long_Plant is not None else "NA",
                r.Long_Shp if r.Long_Shp is not None else "NA",
                r.Long_Correcta if r.Long_Correcta is not None else "NA",
                r.Dec_Plant if r.Dec_Plant is not None else "NA",
                r.Dec_Shp if r.Dec_Shp is not None else "NA",
                r.Dec_Correcto if r.Dec_Correcto is not None else "NA",
                r.orden_correcto,
            ]
            ws.append(row_vals)
            row_idx = ws.max_row
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.font = self.normal_font
                c.border = self.border

            # Colorear celdas booleanas / de estado
            if r.Existe_Campo == "No":
                ws.cell(row=row_idx, column=5).fill = self.fill_err
            else:
                ws.cell(row=row_idx, column=5).fill = self.fill_ok

            for col_i, val_b in [(6, r.Coin_Tipo_Campo), (9, r.Long_Correcta), (12, r.Dec_Correcto), (13, r.orden_correcto)]:
                cell = ws.cell(row=row_idx, column=col_i)
                if val_b is True:
                    cell.fill = self.fill_ok
                elif val_b is False:
                    cell.fill = self.fill_err
                else:
                    cell.fill = self.fill_na

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        self._autofit(ws)

    def _build_sheet_resumen(self, wb: openpyxl.Workbook, result: PAS151AuditResult) -> None:
        ws = wb.create_sheet(title="Resumen_Ejecutivo")
        ws.cell(row=1, column=1, value="CONTROL DE CALIDAD CARTOGRÁFICO PAS 151 - PROTOCOLO CONAF").font = self.title_font
        ws.cell(row=2, column=1, value=f"Fecha de Ejecución: {result.timestamp.strftime('%Y-%m-%d %H:%M:%S')}").font = self.normal_font
        ws.cell(row=3, column=1, value=f"Directorio Auditado: {result.directorio}").font = self.normal_font
        ws.cell(row=4, column=1, value=f"Predio (Sufijo Archivos): {result.nombre_predio}").font = self.bold_font
        ws.cell(
            row=5,
            column=1,
            value=f"Nom_Predio Referencia (Limite_Predial): {result.nom_predio_referencia_atrib or 'No disponible'}",
        ).font = self.bold_font
        ws.cell(row=6, column=1, value=f"Compilado Áreas de Corta: {result.ruta_shp_area_comp or 'No especificado'}").font = self.normal_font

        ws.cell(row=8, column=1, value="Indicador General").font = self.sub_font
        ws.cell(row=8, column=1).fill = self.sub_fill
        ws.cell(row=8, column=2, value="Valor").font = self.sub_font
        ws.cell(row=8, column=2).fill = self.sub_fill

        kpis = [
            ("Total Archivos .shp en Carpeta", len(result.shapes_en_carpeta)),
            ("Total Capas Asociadas y Auditadas", len(result.shapes_auditados)),
            ("Capas Conformes (Sin Errores ni Advertencias)", result.total_capas_conformes),
            ("Capas con Errores", result.total_capas_con_error),
            ("Capas Adicionales (Fuera de Plantilla)", len(result.capas_adicionales)),
            ("Total Incidencias de ERROR", result.total_errores),
            ("Total Incidencias de ADVERTENCIA", result.total_advertencias),
        ]

        for idx, (label, val) in enumerate(kpis, start=9):
            c1 = ws.cell(row=idx, column=1, value=label)
            c2 = ws.cell(row=idx, column=2, value=val)
            c1.font = self.normal_font
            c2.font = self.bold_font
            c1.border = self.border
            c2.border = self.border
            if "Error" in label and val > 0:
                c2.fill = self.fill_err
            elif "Conformes" in label and val > 0:
                c2.fill = self.fill_ok
            elif ("Adicionales" in label or "ADVERTENCIA" in label) and val > 0:
                c2.fill = self.fill_warn

        self._autofit(ws)

    def _build_sheet_validacion_capas(self, wb: openpyxl.Workbook, result: PAS151AuditResult) -> None:
        ws = wb.create_sheet(title="Validacion_Capas")
        headers = [
            "Shapefile",
            "Plantilla CONAF",
            "Registros",
            "CRS (EPSG:32719)",
            "Geometría 2D (Sin Z/M)",
            "Estructura Campos",
            "Orden Campos",
            "Dominios CONAF",
            "Sin Vacíos/NA",
            "Sup_ha vs Geometría",
            "Coord_X / Coord_Y",
            "Texto Sin Acentos",
            "Texto Sin Doble Espacio",
            "Validez Topológica",
            "Estado Global",
        ]
        ws.append(headers)
        self._style_header(ws)

        def fmt_bool(b):
            if b is True:
                return "OK"
            if b is False:
                return "ERROR"
            return "N/A"

        for s in result.resumen_capas:
            row = [
                s.nombre_shape,
                s.prefijo_plantilla,
                s.num_registros,
                "OK" if s.crs_ok else f"ERROR ({s.crs_detalle})",
                "OK" if s.es_2d_ok else f"ERROR ({s.geometria_detalle})",
                fmt_bool(s.estructura_campos_ok),
                fmt_bool(s.orden_campos_ok),
                fmt_bool(s.dominios_ok),
                fmt_bool(s.campos_sin_vacios_ok),
                fmt_bool(s.sup_ha_ok),
                fmt_bool(s.coord_xy_ok),
                fmt_bool(s.texto_caracteres_ok),
                fmt_bool(s.texto_doble_espacio_ok),
                "OK" if s.geometrias_validas_ok else f"ERROR ({s.num_geometrias_invalidas} inv.)",
                s.estado_global,
            ]
            ws.append(row)
            r_idx = ws.max_row
            for col_i in range(1, len(headers) + 1):
                cell = ws.cell(row=r_idx, column=col_i)
                cell.font = self.normal_font
                cell.border = self.border
                val_s = str(cell.value or "")
                if col_i >= 4:
                    if val_s.startswith("OK") or val_s == "CONFORME":
                        cell.fill = self.fill_ok
                    elif "ERROR" in val_s:
                        cell.fill = self.fill_err
                    elif "ADVERTENCIA" in val_s:
                        cell.fill = self.fill_warn
                    elif val_s == "N/A":
                        cell.fill = self.fill_na

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        self._autofit(ws)

    def _build_sheet_espaciales(self, wb: openpyxl.Workbook, result: PAS151AuditResult) -> None:
        ws = wb.create_sheet(title="Validaciones_Espaciales")
        headers = [
            "ID Control",
            "Grupo de Validación Espacial",
            "Capas Evaluadas",
            "Estado",
            "Resultado",
            "Detalle Técnico",
        ]
        ws.append(headers)
        self._style_header(ws)

        for sp in result.resultados_espaciales:
            ws.append([
                sp.control_id,
                sp.grupo,
                sp.capas_evaluadas,
                sp.estado,
                sp.resumen_resultado,
                sp.detalle_tecnico,
            ])
            r_idx = ws.max_row
            for col_i in range(1, len(headers) + 1):
                c = ws.cell(row=r_idx, column=col_i)
                c.font = self.normal_font
                c.border = self.border
            est_cell = ws.cell(row=r_idx, column=4)
            if sp.estado == "CONFORME":
                est_cell.fill = self.fill_ok
            elif sp.estado == "ERROR":
                est_cell.fill = self.fill_err
            elif sp.estado == "ADVERTENCIA":
                est_cell.fill = self.fill_warn
            else:
                est_cell.fill = self.fill_na

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        self._autofit(ws)

    def _build_sheet_incidencias(self, wb: openpyxl.Workbook, result: PAS151AuditResult) -> None:
        ws = wb.create_sheet(title="Incidencias_Detalladas")
        headers = [
            "ID Incidencia",
            "Shapefile",
            "Plantilla",
            "Categoría",
            "Severidad",
            "Campo / Regla",
            "Valor Esperado",
            "Valor Encontrado",
            "Filas Afectadas",
            "Descripción",
            "Recomendación",
        ]
        ws.append(headers)
        self._style_header(ws)

        for iss in result.incidencias:
            ws.append([
                iss.issue_id,
                iss.capa,
                iss.prefijo_plantilla,
                iss.categoria,
                iss.severidad,
                iss.campo_o_regla,
                iss.valor_esperado,
                iss.valor_encontrado,
                iss.filas_afectadas,
                iss.mensaje,
                iss.recomendacion,
            ])
            r_idx = ws.max_row
            for col_i in range(1, len(headers) + 1):
                c = ws.cell(row=r_idx, column=col_i)
                c.font = self.normal_font
                c.border = self.border
            sev_cell = ws.cell(row=r_idx, column=5)
            if iss.severidad == "ERROR":
                sev_cell.fill = self.fill_err
            elif iss.severidad == "ADVERTENCIA":
                sev_cell.fill = self.fill_warn

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        self._autofit(ws)

    def _build_sheet_plantilla(self, wb: openpyxl.Workbook, result: PAS151AuditResult) -> None:
        ws = wb.create_sheet(title="Plantilla_CONAF_PAS151")
        headers = ["shp", "nombre", "tipo", "longitud", "decimales", "orden_esperado", "dominio_conaf"]
        ws.append(headers)
        self._style_header(ws)

        for tmpl in result.plantillas_usadas.values():
            for c in tmpl.campos:
                dom_vals = DOMINIOS_CONAF.get(c.nombre)
                dom_str = ", ".join(str(x) for x in dom_vals) if dom_vals else ""
                ws.append([
                    tmpl.prefijo,
                    c.nombre,
                    c.tipo,
                    c.longitud if c.longitud is not None else "NA",
                    c.decimales if c.decimales is not None else "NA",
                    c.orden,
                    dom_str,
                ])
                r_idx = ws.max_row
                for col_i in range(1, len(headers) + 1):
                    cell = ws.cell(row=r_idx, column=col_i)
                    cell.font = self.normal_font
                    cell.border = self.border

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        self._autofit(ws)

    def _style_header(self, ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
        for cell in ws[1]:
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _autofit(self, ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 55)
