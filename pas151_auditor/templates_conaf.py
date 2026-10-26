# -*- coding: utf-8 -*-
"""
Plantillas oficiales de capas, atributos y dominios del Protocolo CONAF para PAS 151 (Rev 11),
más soporte para carga dinámica desde archivo Excel (Diccionario_PAS151.xlsx) respetando la Regla NA.
"""

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union
import unicodedata


@dataclass
class FieldSpec:
    """Especificación de un campo según la plantilla CONAF PAS 151."""
    nombre: str
    tipo: str  # 'character', 'integer', 'numeric'
    longitud: Optional[int] = None   # None equivale a NA en el script R
    decimales: Optional[int] = None  # None equivale a NA en el script R
    orden: int = 0


@dataclass
class LayerTemplate:
    """Plantilla estructural de un shapefile del protocolo CONAF."""
    prefijo: str
    campos: List[FieldSpec] = field(default_factory=list)
    origen: str = "PLANTILLA_R_REV11"

    @property
    def nombres_campos(self) -> List[str]:
        return [c.nombre for c in self.campos]

    def obtener_campo(self, nombre: str) -> Optional[FieldSpec]:
        nombre_up = nombre.strip().upper()
        for c in self.campos:
            if c.nombre.strip().upper() == nombre_up:
                return c
        return None


def normalize_layer_key(text: str) -> str:
    """
    Normaliza nombres de prefijos de capas para tolerar diferencias de codificación
    (ej. 'Señaletica_velocidad' vs 'Sealetica_velocidad' o 'Senaletica_velocidad').
    """
    if not text:
        return ""
    s = str(text).strip().replace("\ufffd", "")
    # Reemplazar ñ/Ñ por n/N antes de quitar diacríticos
    s = s.replace("ñ", "n").replace("Ñ", "N")
    s = "".join(
        c for c in unicodedata.normalize("NFD", s)
        if unicodedata.category(c) != "Mn"
    )
    # También normalizar caso donde 'ñ' se perdió como caracter inválido ('Sealetica')
    s_up = s.upper()
    if s_up.startswith("SEALETICA"):
        s_up = "SENALETICA" + s_up[len("SEALETICA"):]
    return s_up


# Definición exacta de las plantillas de Revision_Campos_PAS151_Rev11 2.R + Diccionario_PAS151.xlsx
RAW_TEMPLATES_PAS151: Dict[str, List[Tuple[str, str, Optional[int], Optional[int]]]] = {
    "Area": [
        ("Nom_Predio", "character", None, None),
        ("N_Area", "character", 5, None),
        ("Tipo_Bos", "character", 10, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Limite_Predial": [
        ("Nom_Predio", "character", None, None),
        ("Rol", "character", 50, None),
        ("Comuna", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Suelos": [
        ("Nom_Predio", "character", None, None),
        ("Cat_Erosio", "character", 2, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Rangos_pend": [
        ("Nom_Predio", "character", None, None),
        ("Ran_Pend", "character", 10, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Infra_aislada": [
        ("Nom_Predio", "character", None, None),
        ("Nom_Infra", "character", 50, None),
        ("Coord_X", "integer", 6, 0),
        ("Coord_Y", "integer", 7, 0),
        ("Fuente", "character", None, None),
    ],
    "Parcela": [
        ("Nom_Predio", "character", None, None),
        ("N_Rodal", "character", 5, 0),
        ("N_Parc", "character", 5, 0),
        ("Coord_X", "integer", 6, 0),
        ("Coord_Y", "integer", 7, 0),
        ("Fuente", "character", None, None),
    ],
    "Referencia": [
        ("Nom_Predio", "character", None, None),
        ("Nom_Pto", "character", 100, None),
        ("Coord_X", "integer", 6, 0),
        ("Coord_Y", "integer", 7, 0),
        ("Fuente", "character", None, None),
    ],
    "Caminos": [
        ("Nom_Predio", "character", None, None),
        ("Tipo_Cam", "character", 2, None),
        ("Long_Cam", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Curvas_niv": [
        ("Nom_Predio", "character", None, None),
        ("Cot_Curva", "integer", 9, None),
        ("Long_Curva", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Hidrografia": [
        ("Nom_Predio", "character", None, None),
        ("Tip_Dren", "character", 2, None),
        ("Tipo_Perma", "character", 2, None),
        ("Long_Dren", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Rodales": [
        ("Nom_Predio", "character", None, None),
        ("N_Rodal", "character", 5, None),
        ("Tipo_Bos", "character", 10, None),
        ("Tipo_For", "character", 10, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Predios_vecinos": [
        ("Nom_Predio", "character", None, None),
        ("Rol", "character", 50, None),
        ("Comuna", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "LTE_HVDC": [
        ("Nom_Predio", "character", None, None),
        ("Obra", "character", 80, None),
        ("Nom_obra", "character", 80, None),
        ("Obra_detal", "character", 80, None),
        ("Tipo", "character", 80, None),
        ("Tramo", "character", 80, None),
        ("Long_LTE", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Sendero_peatonal": [
        ("Nom_Predio", "character", None, None),
        ("Nom_Send", "character", 50, None),
        ("Long_Send", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Area_Proyecto": [
        ("Nom_Predio", "character", None, None),
        ("Nom_obra", "character", 80, None),
        ("Obra_detal", "character", 80, None),
        ("Tipo", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Area_proteccion": [
        ("Nom_Predio", "character", None, None),
        ("Tipo", "character", 100, None),
        ("Nombre", "character", 100, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Limite_Comunal": [
        ("Region", "character", 50, None),
        ("Provincia", "character", 50, None),
        ("Comuna", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Servidumbre": [
        ("Nom_Predio", "character", None, None),
        ("Rol", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    # MEDIDAS FAUNA
    "Areas_receptoras_ruido": [
        ("Nom_Predio", "character", None, None),
        ("Nom_obra", "character", 50, None),
        ("Obra_detal", "character", 50, None),
        ("Ambiente", "character", 50, None),
        ("Receptor", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Perturbacion_ruido": [
        ("Nom_Predio", "character", None, None),
        ("Nom_obra", "character", 50, None),
        ("Obra_detal", "character", 50, None),
        ("Ambiente", "character", 50, None),
        ("Receptor", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Perturbacion_baja_mov": [
        ("Nom_Predio", "character", None, None),
        ("Nom_obra", "character", 50, None),
        ("Obra_detal", "character", 50, None),
        ("Perturbaci", "character", 50, None),
        ("Ambiente", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Receptoras_baja_mov": [
        ("Nom_Predio", "character", None, None),
        ("Nom_obra", "character", 50, None),
        ("Obra_detal", "character", 50, None),
        ("Perturbaci", "character", 50, None),
        ("Ambiente", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Perturbacion_curureras": [
        ("Nom_Predio", "character", None, None),
        ("Obra_detal", "character", 50, None),
        ("Ambiente", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Receptoras_curureras": [
        ("Nom_Predio", "character", None, None),
        ("Obra_detal", "character", 50, None),
        ("Ambiente", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Perchas": [
        ("Nom_Predio", "character", None, None),
        ("Medida", "character", 50, None),
        ("Coord_X", "integer", 6, 0),
        ("Coord_Y", "integer", 7, 0),
        ("Fuente", "character", None, None),
    ],
    "Señaletica_velocidad": [
        ("Nom_Predio", "character", None, None),
        ("Tipo_inst", "character", 50, None),
        ("Coord_X", "integer", 6, 0),
        ("Coord_Y", "integer", 7, 0),
        ("Fuente", "character", None, None),
    ],
    # MEDIDAS PLANTAS SUELO
    "Rehabilitacion": [
        ("Nom_Predio", "character", None, None),
        ("Nom_obra", "character", 50, None),
        ("Obra_detal", "character", 50, None),
        ("Medida", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Surcos": [
        ("Nom_Predio", "character", None, None),
        ("Medida", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Rehabilitacion_pendiente_30": [
        ("Nom_Predio", "character", None, None),
        ("Nom_obra", "character", 50, None),
        ("Obra_detal", "character", 50, None),
        ("Medida", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Recoleccion_semillas": [
        ("Nom_Predio", "character", None, None),
        ("Medida", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Sectores_recoleccion_semillas": [
        ("Nom_Predio", "character", None, None),
        ("Nom_obra", "character", 10, None),
        ("Obra_detal", "character", 10, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Sectores_siembra_directa": [
        ("Nom_Predio", "character", None, None),
        ("Tipo_Veg", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    # MEDIDAS PLANTAS
    "Areas_relocalizacion": [
        ("Nom_Predio", "character", None, None),
        ("Medida", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Areas_rescate": [
        ("Nom_Predio", "character", None, None),
        ("Nom_obra", "character", 50, None),
        ("Obra_detal", "character", 50, None),
        ("Especie", "character", 50, None),
        ("Medida", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Reposicion": [
        ("Nom_Predio", "character", None, None),
        ("Medida", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    # REVEGETACION
    "Revegetacion": [
        ("Nom_Predio", "character", None, None),
        ("Nom_Sitio", "character", None, None),
        ("Region", "character", 50, None),
        ("Provincia", "character", 50, None),
        ("Comuna", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Area_FX_singular": [
        ("Nom_Predio", "character", None, None),
        ("Nom_obra", "character", 50, None),
        ("Obra_detal", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    # ENRIQUECIMIENTO
    "Enriquecimiento": [
        ("Nom_Predio", "character", None, None),
        ("Nom_Sitio", "character", None, None),
        ("Region", "character", 50, None),
        ("Provincia", "character", 50, None),
        ("Comuna", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    "Areas_ECC": [
        ("Nom_Predio", "character", None, None),
        ("Nom_obra", "character", None, None),
        ("Obra_detal", "character", 50, None),
        ("Especie", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
    # MEDIDAS ADICIONALES
    "Infra_incendio": [
        ("Nom_Predio", "character", None, None),
        ("Nom_Infra", "character", 50, None),
        ("Coord_X", "integer", 6, 0),
        ("Coord_Y", "integer", 7, 0),
        ("Fuente", "character", None, None),
    ],
    "Obras_conservacion": [
        ("Nom_Predio", "character", None, None),
        ("Nom_Infra", "character", 50, None),
        ("Coord_X", "integer", 6, 0),
        ("Coord_Y", "integer", 7, 0),
        ("Fuente", "character", None, None),
    ],
    "Badenes": [
        ("Nom_Predio", "character", None, None),
        ("Nom_Infra", "character", 50, None),
        ("Coord_X", "integer", 6, 0),
        ("Coord_Y", "integer", 7, 0),
        ("Fuente", "character", None, None),
    ],
    "Pendientes_caminos": [
        ("Nom_Predio", "character", None, None),
        ("Nom_obra", "character", 80, None),
        ("Obra_detal", "character", 80, None),
        ("Ran_Pend", "character", 50, None),
        ("Medida", "character", 50, None),
        ("Sup_ha", "numeric", 11, 2),
        ("Fuente", "character", None, None),
    ],
}

# Dominios del protocolo CONAF definidos en Revision_Campos_PAS151_Rev11 2.R
DOMINIOS_CONAF: Dict[str, List[str]] = {
    "Tipo_Bos": ["No aplica", "Bosque nativo", "Plantaciones"],
    "Tipo_For": ["No aplica"] + [str(i) for i in range(1, 13)],
    "Tipo_Cam": ["1", "2", "3", "4"],
    "Tip_Dren": ["1", "2", "3", "4", "5"],
    "Tipo_Perma": ["1", "2"],
    "Cat_Erosio": ["1", "2", "3", "4"],
    "Ran_Pend": ["0% - 10%", "10% - 30%", "30% - 45%", "45% - 60%", "60% y mas"],
}

# Categorías de Infra_aislada que no deben ubicarse sobre Area_Proyecto (líneas 817-825 del R)
CATEGORIAS_INFRA_EROSION: List[str] = [
    "Control de erosion",
    "Sacos de tierra",
    "Bordo de piedra",
    "Zanja infiltracion",
    "Empalizada",
    "Fajina",
    "Microterraza",
]

# Capas cuya contención dentro de Limite_Predial y Area_Proyecto se valida (línea 737 del R + Descripcion_general.docx)
CAPAS_CONTENCION_PREDIO: List[str] = [
    "Area",
    "Suelos",
    "Rangos_pend",
    "Pendientes_caminos",
]

# Capas excluidas de la validación transversal de Nom_Predio respecto de Limite_Predial
CAPAS_EXCLUIDAS_NOM_PREDIO: List[str] = [
    "Predios_vecinos",
    "Limite_Comunal",
]


def _parse_optional_int(val: Any) -> Optional[int]:
    """Aplica la Regla NA para longitud y decimales."""
    if val is None:
        return None
    s = str(val).strip()
    if not s or s.upper() in ("NA", "N/A", "NONE", "NULL", "NAN"):
        return None
    try:
        return int(float(s))
    except (ValueError, TypeError):
        return None


def get_builtin_templates() -> Dict[str, LayerTemplate]:
    """Retorna las plantillas oficiales PAS 151 Rev11 indexadas por clave normalizada."""
    result: Dict[str, LayerTemplate] = {}
    for prefijo, fields_raw in RAW_TEMPLATES_PAS151.items():
        campos = [
            FieldSpec(
                nombre=f[0],
                tipo=f[1].lower(),
                longitud=f[2],
                decimales=f[3],
                orden=idx + 1,
            )
            for idx, f in enumerate(fields_raw)
        ]
        key = normalize_layer_key(prefijo)
        result[key] = LayerTemplate(prefijo=prefijo, campos=campos, origen="PLANTILLA_R_REV11")
    return result


def load_templates(excel_path: Optional[Union[str, Path]] = None) -> Dict[str, LayerTemplate]:
    """
    Carga las plantillas oficiales PAS 151 Rev11 y, si se provee un archivo Excel
    válido (ej. Diccionario_PAS151.xlsx), actualiza o agrega las definiciones desde el Excel.
    """
    templates = get_builtin_templates()
    if not excel_path:
        return templates

    path_obj = Path(excel_path)
    if not path_obj.exists():
        return templates

    import openpyxl

    wb = openpyxl.load_workbook(path_obj, data_only=True)
    target_ws = wb.active
    required = {"shp", "nombre", "tipo", "longitud", "decimales"}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if first_row:
            cols = {str(c).strip().lower() for c in first_row if c is not None}
            if required.issubset(cols):
                target_ws = ws
                break

    rows = list(target_ws.iter_rows(values_only=True))
    if not rows:
        return templates

    headers = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    try:
        idx_shp = headers.index("shp")
        idx_nom = headers.index("nombre")
        idx_tip = headers.index("tipo")
        idx_lon = headers.index("longitud")
        idx_dec = headers.index("decimales")
    except ValueError:
        return templates

    excel_grouped: Dict[str, Tuple[str, List[FieldSpec]]] = {}
    for row in rows[1:]:
        if not row or len(row) <= max(idx_shp, idx_nom, idx_tip):
            continue
        raw_shp = row[idx_shp]
        raw_nom = row[idx_nom]
        raw_tip = row[idx_tip]
        if raw_shp is None or raw_nom is None or raw_tip is None:
            continue

        shp_str = str(raw_shp).strip()
        nom_str = str(raw_nom).strip()
        tip_str = str(raw_tip).strip().lower()
        if not shp_str or not nom_str:
            continue

        # Restaurar 'Señaletica_velocidad' si el Excel tenía caracter corrupto
        if normalize_layer_key(shp_str) == "SENALETICA_VELOCIDAD":
            shp_str = "Señaletica_velocidad"

        key = normalize_layer_key(shp_str)
        lon_val = _parse_optional_int(row[idx_lon] if idx_lon < len(row) else None)
        dec_val = _parse_optional_int(row[idx_dec] if idx_dec < len(row) else None)

        if key not in excel_grouped:
            excel_grouped[key] = (shp_str, [])
        orden = len(excel_grouped[key][1]) + 1
        excel_grouped[key][1].append(
            FieldSpec(
                nombre=nom_str,
                tipo=tip_str,
                longitud=lon_val,
                decimales=dec_val,
                orden=orden,
            )
        )

    for key, (prefijo_str, campos_list) in excel_grouped.items():
        templates[key] = LayerTemplate(
            prefijo=prefijo_str,
            campos=campos_list,
            origen=f"EXCEL ({path_obj.name})",
        )

    return templates
